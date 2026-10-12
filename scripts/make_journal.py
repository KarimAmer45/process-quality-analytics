"""Author with artifact-tool when installed, otherwise use the portable openpyxl writer."""
import json,os,pathlib,shutil,subprocess
from datetime import date

ROOT=pathlib.Path(__file__).resolve().parents[1]

def main():
    node=os.environ.get('NODE_BINARY') or shutil.which('node')
    if node and (ROOT/'node_modules/@oai/artifact-tool').exists():
        subprocess.run([node,str(ROOT/'scripts/build_journal.mjs')],cwd=ROOT,check=True)
        return
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active)
    payload=json.loads((ROOT/'data/interim/journal.json').read_text(encoding='utf-8'))
    for name,rows in payload.items():
        ws=wb.create_sheet(name);width=9 if name=='Analysen' else 7
        for row in rows:
            ws.append([date.fromisoformat(v['date']) if isinstance(v,dict) and 'date' in v else v for v in row])
        for row in ws:
            for cell in row:
                cell.font=Font(name='Arial',size=10)
                if isinstance(cell.value,date):cell.number_format='yyyy-mm-dd'
            ws.row_dimensions[row[0].row].height=21
        for r in (1,2):ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=width)
        ws['A1'].font=Font(name='Arial',size=15,bold=True)
        ws['A2'].font=Font(name='Arial',size=10,italic=True,color='526579')
        for cell in ws[4]:cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='233B57')
        for c in range(1,width+1):ws.column_dimensions[get_column_letter(c)].width=18
        ws.column_dimensions[get_column_letter(width)].width=45
        ws.freeze_panes='A5';ws.sheet_view.showGridLines=False
    out=ROOT/'data/raw/lab/Laborjournal_ZnNi_L1.xlsx';out.parent.mkdir(parents=True,exist_ok=True)
    wb.save(out);print(f'Portable openpyxl workbook: {out.relative_to(ROOT).as_posix()}')

if __name__=='__main__':main()
