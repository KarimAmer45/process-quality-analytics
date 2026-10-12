from datetime import date,time
from decimal import Decimal
import pytest
from openpyxl import Workbook,load_workbook
from etl_lab_journal import RowError,parse_number,parse_date,parse_time,local_timestamp,read_records,import_workbook

@pytest.mark.parametrize('raw,expected',[('1,25',Decimal('1.25')),('1.25',Decimal('1.25')),(0,Decimal(0)),('0',Decimal(0)),('n.b.',None),(None,None)])
def test_numbers(raw,expected): assert parse_number(raw)==expected

@pytest.mark.parametrize('raw',['1,3.5','NaN',float('inf'),True,-1,'=1+1'])
def test_bad_numbers(raw):
    with pytest.raises(RowError):parse_number(raw)

@pytest.mark.parametrize('raw',['03.03.2026','3/3/26','2026-03-03'])
def test_dates(raw): assert parse_date(raw)==date(2026,3,3)

def test_times():
    assert parse_time(None)==(None,'time missing')   # never imputed
    assert parse_time(.25)[0]==time(6)
    with pytest.raises(RowError):parse_time('25:99')

@pytest.mark.parametrize('day',[date(2026,3,29),date(2026,10,25)])
def test_dst_ambiguity_rejected(day):
    with pytest.raises(RowError):local_timestamp(day,time(2,30))

def fixture_book(path,rows=None,doses=None):
    # Small disposable test fixture, not a deliverable workbook. Dates lie before the simulated
    # period (and not after the import day, which the date plausibility window rejects).
    wb=Workbook();ws=wb.active;ws.title='Analysen'
    ws.append(['Datum','Uhrzeit','Bad','Zn [g/l]','Ni [g/l]','NaOH [g/l]','Na2CO3 [g/l]','CSB [g/l]'])
    for row in rows or [['01.01.2025','06:00','L1-T05','10','1,4','120','n.b.','100']]:ws.append(row)
    ds=wb.create_sheet('Dosierung');ds.append(['Datum','Uhrzeit','Bad','Zusatz','Menge','Einheit'])
    for dose in doses or [['01.01.2025','12:00','L1-T05','brightener','1,5','L']]:ds.append(dose)
    wb.save(path)

def test_conflicting_duplicates_both_quarantined(tmp_path):
    path=tmp_path/'conflict.xlsx'
    row=['01.01.2025','06:00','L1-T05','10','1.4','120','70','100']
    other=list(row);other[4]='1.5'
    fixture_book(path,[row,other]);parsed=read_records(path)
    assert sum(r['status']=='quarantined' for r in parsed)==2

def test_dry_run_changes_nothing_and_apply_is_idempotent(db,tmp_path):
    path=tmp_path/'input.xlsx';fixture_book(path)
    before=db.execute('SELECT count(*) FROM lab.analysis').fetchone()[0]
    dry=import_workbook(db,path)
    assert not dry['applied']
    assert db.execute('SELECT count(*) FROM lab.analysis').fetchone()[0]==before
    assert not db.execute('SELECT 1 FROM lab.import_file WHERE sha256=%s',(dry['sha256'],)).fetchone()
    first=import_workbook(db,path,True)
    snapshot=db.execute('SELECT sha256,applied_at,report FROM lab.import_file ORDER BY sha256').fetchall()
    second=import_workbook(db,path,True)
    assert first['inserted']=={'analyses':1,'results':4,'doses':1}
    assert second['reused'] and second['inserted_this_run']==0
    assert snapshot==db.execute('SELECT sha256,applied_at,report FROM lab.import_file ORDER BY sha256').fetchall()
    assert db.execute("SELECT amount FROM lab.dosing WHERE dosed_at='2025-01-01T12:00:00+01:00' AND additive='brightener'").fetchone()[0]==1500

def test_existing_value_conflict_not_overwritten(db,tmp_path):
    path=tmp_path/'input.xlsx';fixture_book(path)
    import_workbook(db,path,True)
    changed=tmp_path/'changed.xlsx'
    fixture_book(changed,[['01.01.2025','06:00','L1-T05','10','1.6','120','n.b.','100']])
    result=import_workbook(db,changed,True)
    assert result['reconciliation_before']['conflicting_rows']==1
    assert db.execute("SELECT r.value FROM lab.analysis a JOIN lab.analysis_result r USING(analysis_id) WHERE sampled_at='2025-01-01T06:00:00+01:00' AND parameter_code='NI'").fetchone()[0]==Decimal('1.4')

def test_unreadable_row_of_a_stored_sample_is_one_finding(db,tmp_path):
    # The row is still in the workbook: it is quarantined, and its stored sample and dose are not also reported as missing.
    rows=[[f'0{d}.01.2025','06:00','L1-T05','10','1,4','120','n.b.','100'] for d in (1,2,3)]
    doses=[[f'0{d}.01.2025','12:00','L1-T05','brightener','1,5','L'] for d in (1,2,3)]
    path=tmp_path/'input.xlsx';fixture_book(path,rows,doses);import_workbook(db,path,True)
    rows[1][4]='1,3.5';doses[1][4]='viel'
    changed=tmp_path/'typo.xlsx';fixture_book(changed,rows,doses)
    report=import_workbook(db,changed)
    assert report['review_required']=={'quarantined_rows':2,'database_only_rows':0,'retimed_daily_totals':0}
    assert report['database_only_rows']==[] and not any(i['status']=='database_only' for i in report['issues'])

def test_equal_filenames_different_content_keep_history(db,tmp_path):
    path=tmp_path/'journal.xlsx';fixture_book(path)
    first=import_workbook(db,path,True)
    fixture_book(path,[['02.01.2025','06:00','L1-T05','10','1.4','120','n.b.','100']])
    second=import_workbook(db,path,True)
    assert first['sha256']!=second['sha256']
    assert db.execute('SELECT count(*) FROM lab.import_file WHERE sha256=ANY(%s)',([first['sha256'],second['sha256']],)).fetchone()[0]==2

def test_wrong_concentration_units_rejected(tmp_path):
    path=tmp_path/'units.xlsx';fixture_book(path)
    wb=load_workbook(path);wb['Analysen']['D1']='Zn [mg/l]';wb.save(path)
    with pytest.raises(ValueError,match='expected g/L'):read_records(path)

def test_extra_sheet_requires_review(tmp_path):
    path=tmp_path/'extra.xlsx';fixture_book(path)
    wb=load_workbook(path);wb.create_sheet('Unreviewed');wb.save(path)
    with pytest.raises(ValueError,match='Unrecognized sheets'):read_records(path)

def test_dosing_precision_is_not_silently_rounded(tmp_path):
    path=tmp_path/'dose.xlsx';fixture_book(path)
    wb=load_workbook(path);wb['Dosierung']['E2']='0.0000001';wb.save(path)
    records=read_records(path)
    assert any(r['status']=='quarantined' and 'precision' in r['reason'] for r in records)

# Reconciliation in both directions, re-keyed rows, daily-total identity, dates and summaries.
def row(day,clock='06:00',ni='1,4'):return [day,clock,'L1-T05','10',ni,'120','70','100']
def dose(day,clock='12:00',amount='1,5'):return [day,clock,'L1-T05','brightener',amount,'L']
def count(db,sql,*args):return db.execute(sql,args).fetchone()[0]

def test_rows_missing_from_a_new_version_are_reported(db,tmp_path):
    first=tmp_path/'a'/'journal.xlsx';first.parent.mkdir()
    fixture_book(first,[row('01.01.2025'),row('02.01.2025'),row('03.01.2025')],[dose('01.01.2025'),dose('02.01.2025'),dose('03.01.2025')])
    assert import_workbook(db,first,True)['review_required']['database_only_rows']==0
    second=tmp_path/'b'/'journal.xlsx';second.parent.mkdir()
    fixture_book(second,[row('01.01.2025'),row('03.01.2025')],[dose('01.01.2025'),dose('03.01.2025')])
    report=import_workbook(db,second,True)
    assert report['reconciliation_before']['database_only_analyses']==1
    assert report['reconciliation_before']['database_only_doses']==1
    assert report['review_required']['database_only_rows']==2 and report['inserted_this_run']==0
    assert sum(i['status']=='database_only' for i in report['issues'])==2
    assert report['database_only_rows'][0]['sampled_at'].startswith('2025-01-02T06:00')

def test_changed_time_is_quarantined_not_inserted(db,tmp_path):
    first=tmp_path/'a'/'journal.xlsx';first.parent.mkdir();fixture_book(first,[row('01.01.2025')])
    import_workbook(db,first,True)
    second=tmp_path/'b'/'journal.xlsx';second.parent.mkdir();fixture_book(second,[row('01.01.2025','06:30')])
    report=import_workbook(db,second,True)
    assert report['reconciliation_before']['rekeyed_rows']==1 and report['statuses']['quarantined']==1
    assert report['inserted']['analyses']==0 and report['review_required']['database_only_rows']==0
    assert count(db,"SELECT count(*) FROM lab.analysis WHERE sampled_at>='2025-01-01' AND sampled_at<'2025-01-02'")==1

def test_second_sample_on_a_day_is_new_data_not_a_rekey(db,tmp_path):
    first=tmp_path/'a'/'journal.xlsx';first.parent.mkdir();fixture_book(first,[row('01.01.2025')])
    import_workbook(db,first,True)
    second=tmp_path/'b'/'journal.xlsx';second.parent.mkdir();fixture_book(second,[row('01.01.2025'),row('01.01.2025','14:00',ni='1,5')])
    report=import_workbook(db,second,True)
    assert report['reconciliation_before']['rekeyed_rows']==0 and report['inserted']['analyses']==1

def test_daily_total_identity_is_tank_berlin_day_and_additive(db,tmp_path):
    def run(name,doses):
        path=tmp_path/name/'journal.xlsx';path.parent.mkdir();fixture_book(path,None,doses);return import_workbook(db,path,True)
    run('a',[dose('01.01.2025')])
    moved=run('b',[dose('01.01.2025','13:00')])
    assert moved['inserted']['doses']==0 and moved['reconciliation_before']['retimed_daily_totals']==1
    no_time=run('c',[dose('01.01.2025',None)])
    assert no_time['inserted']['doses']==0 and no_time['statuses'].get('quarantined',0)==0
    changed=run('d',[dose('01.01.2025','12:00','1,6')])
    assert changed['reconciliation_before']['conflicting_rows']==1 and changed['inserted']['doses']==0
    new_without_time=run('e',[dose('01.01.2025'),dose('02.01.2025',None)])
    assert new_without_time['statuses']['quarantined']==1 and new_without_time['inserted']['doses']==0
    assert count(db,"SELECT count(*) FROM lab.dosing WHERE dosed_at>='2025-01-01' AND dosed_at<'2025-01-03'")==1

def test_implausible_dates_are_quarantined(tmp_path):
    path=tmp_path/'dates.xlsx';fixture_book(path,[row('03.03.2062'),row('03.03.1962')])
    records=read_records(path)
    assert [r['status'] for r in records if r['sheet']=='Analysen' and r['row']>1]==['quarantined','quarantined']
    assert all('outside plausible window' in r['reason'] for r in records if r['status']=='quarantined')
    wide=read_records(path,(date(1900,1,1),date(2100,1,1)))
    assert wide[2]['status']=='loaded' and 'later than now' in wide[1]['reason']

def test_missing_time_only_confirms_the_existing_sample_of_that_day(db,tmp_path):
    tank=count(db,"SELECT tank_id FROM ref.tank WHERE code='L1-T05'")
    aid=count(db,"INSERT INTO lab.analysis(tank_id,sampled_at,source) VALUES(%s,'2025-01-01 14:00+01','test') RETURNING analysis_id",tank)
    for param,value,method in (('ZN',10,'titration'),('NI',1.4,'AAS'),('NAOH',120,'titration'),('NA2CO3',70,'titration'),('COD',100,'COD')):
        db.execute("INSERT INTO lab.analysis_result VALUES(%s,%s,%s,%s)",(aid,param,value,method))
    path=tmp_path/'time.xlsx';fixture_book(path,[row('01.01.2025',None),row('02.01.2025',None)])
    report=import_workbook(db,path,True)
    reasons={i['row']:i['reason'] for i in report['issues'] if i['sheet']=='Analysen'}
    assert 'matched the only database sample of that date (14:00)' in reasons[2]
    assert 'no database sample on this date' in reasons[3]
    assert report['inserted']['analyses']==0 and report['reconciliation_before']['matched_values']==5
    assert count(db,"SELECT count(*) FROM lab.analysis WHERE source<>'test' AND sampled_at>='2025-01-01' AND sampled_at<'2025-01-03'")==0

def test_summary_label_with_data_and_headerless_values_are_quarantined(tmp_path):
    path=tmp_path/'summary.xlsx'
    fixture_book(path,[['Mittelwert KW 1',None,None,'10','1,4','120','70','100'],
                       ['Mittelwert-Korrektur 01.01.2025','06:00','L1-T05','10','1,4','120','70','100'],
                       row('02.01.2025')+[None,'1,9']])
    status={r['row']:(r['status'],r['reason']) for r in read_records(path) if r['sheet']=='Analysen'}
    assert status[2]==('ignored','weekly summary; not a measurement')
    assert status[3][0]=='quarantined' and 'summary label' in status[3][1]
    assert status[4]==('quarantined','value in a column without header: J')

def test_changed_date_with_unchanged_values_is_quarantined_not_inserted(db,tmp_path):
    first=tmp_path/'a'/'journal.xlsx';first.parent.mkdir();fixture_book(first,[row('01.01.2025'),row('02.01.2025',ni='1,5')])
    import_workbook(db,first,True)
    second=tmp_path/'b'/'journal.xlsx';second.parent.mkdir();fixture_book(second,[row('01.01.2025'),row('04.01.2025',ni='1,5')])
    report=import_workbook(db,second,True)
    assert report['reconciliation_before']['rekeyed_rows']==1 and report['inserted']['analyses']==0
    assert report['reconciliation_before']['source_only_values']==0 and report['review_required']['database_only_rows']==0
    assert any('changed date?' in (i['reason'] or '') for i in report['issues'])

def test_same_amount_at_two_explicit_times_is_not_one_dose(db,tmp_path):
    first=tmp_path/'a'/'journal.xlsx';first.parent.mkdir();fixture_book(first,None,[dose('01.01.2025')])
    import_workbook(db,first,True)
    second=tmp_path/'b'/'journal.xlsx';second.parent.mkdir()
    fixture_book(second,None,[dose('01.01.2025','12:00'),dose('01.01.2025','16:00')])
    report=import_workbook(db,second,True)
    assert report['statuses']['quarantined']==2 and report['inserted']['doses']==0
    assert report['reconciliation_before']['database_only_doses']==1
    # a missing time still confirms the same daily total
    path=tmp_path/'same.xlsx';fixture_book(path,None,[dose('01.01.2025','12:00'),dose('01.01.2025',None)])
    assert [r['status'] for r in read_records(path) if r['sheet']=='Dosierung' and r['row']>1]==['loaded','duplicate']

def test_identical_values_at_a_new_time_and_future_times_are_quarantined(db,tmp_path):
    first=tmp_path/'a'/'journal.xlsx';first.parent.mkdir();fixture_book(first,[row('01.01.2025')])
    import_workbook(db,first,True)
    # 06:00 confirmed with COD not measured; 14:00 repeats the stored 06:00 values
    second=tmp_path/'b'/'journal.xlsx';second.parent.mkdir()
    fixture_book(second,[row('01.01.2025')[:-1]+['n.b.'],row('01.01.2025','14:00')])
    report=import_workbook(db,second,True)
    assert report['statuses']['quarantined']==1 and report['inserted']['analyses']==0
    assert report['reconciliation_before']['rekeyed_rows']==1
    copies=tmp_path/'copies.xlsx';fixture_book(copies,[row('02.01.2025'),row('02.01.2025','14:00')])
    assert sum(r['status']=='quarantined' for r in read_records(copies))==2
    from datetime import datetime,timedelta
    from etl_lab_journal import TZ
    soon=datetime.now(TZ)+timedelta(minutes=5)
    future=tmp_path/'future.xlsx';fixture_book(future,[row(soon.strftime('%d.%m.%Y'),soon.strftime('%H:%M'))])
    records=read_records(future,(date(2020,1,1),soon.date()))
    assert records[1]['status']=='quarantined' and 'later than now' in records[1]['reason']

def test_reused_file_reports_the_current_database_comparison(db,tmp_path):
    path=tmp_path/'journal.xlsx';fixture_book(path,[row('01.01.2025'),row('02.01.2025'),row('03.01.2025',ni='1,5')])
    first=import_workbook(db,path,True)
    assert first['review_required']['database_only_rows']==0
    tank=count(db,"SELECT tank_id FROM ref.tank WHERE code='L1-T05'")
    db.execute("INSERT INTO lab.analysis(tank_id,sampled_at,source) VALUES(%s,'2025-01-02 14:00+01','test')",(tank,))
    again=import_workbook(db,path,True)
    assert again['reused'] and again['inserted_this_run']==0 and again['review_required']['database_only_rows']==0
    assert again['review_required_now']['database_only_rows']==1
    assert again['database_only_rows_now'][0]['sampled_at'].startswith('2025-01-02T14:00')
