"""Loss-aware Excel import. Dry runs never write; the import ledger is insert-only.

Every sheet row is retained with a status. Conflicting duplicates and existing-value
conflicts are quarantined, never silently overwritten. Reconciliation runs in both
directions: workbook rows are matched against the database, and database rows inside the
workbook's tanks and Berlin date span that have no workbook row are reported. Dosing rows
are daily totals identified by tank | Berlin date | additive.
"""
import argparse
import hashlib
import json
import math
import re
from collections import Counter,defaultdict
from itertools import combinations
from datetime import date,datetime,time,timedelta,timezone
from decimal import Decimal,InvalidOperation
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from psycopg.types.json import Jsonb
from common import ROOT,connect,sources

TZ=ZoneInfo('Europe/Berlin')
PARAMS={'zn':'ZN','ni':'NI','naoh':'NAOH','na2co3':'NA2CO3','csb':'COD'}
ALIASES={x:'L1-T05' for x in ('l1-t05','znni 1','zn-ni 1','zink-nickel l1','zinknickel l1')}
METHODS={'ZN':'titration','NI':'AAS','NAOH':'titration','NA2CO3':'titration','COD':'COD'}

class RowError(ValueError): pass

def key(v): return re.sub(r'\s+',' ',str('' if v is None else v).strip().lower())

def parse_number(value):
    if value is None or key(value) in ('n.b.','nb','-',''): return None
    if isinstance(value,bool): raise RowError('boolean is not a measurement')
    text=str(value).strip()
    if not re.fullmatch(r'[+-]?(?:\d+(?:[.,]\d*)?|[.,]\d+)(?:[eE][+-]?\d+)?',text):
        raise RowError(f'not a number: {text!r}')
    try: number=Decimal(text.replace(',','.'))
    except InvalidOperation as exc: raise RowError('invalid number') from exc
    if not number.is_finite() or abs(number)>=Decimal('100000000'): raise RowError('non-finite or out-of-storage-range number')
    if number<0: raise RowError('negative concentration or dose')
    return number

def parse_date(value):
    if isinstance(value,datetime): return value.date()
    if isinstance(value,date): return value
    for fmt in ('%d.%m.%Y','%d/%m/%y','%d/%m/%Y','%Y-%m-%d'):
        try: return datetime.strptime(str(value).strip(),fmt).date()
        except ValueError: pass
    raise RowError(f'not a date: {value!r}')

def date_window(today=None):
    """Plausible sampling/dosing dates from config/sources.yaml: [earliest, import day + allowance]."""
    rule=sources()['assumptions']['lab_journal_dates']['value']
    today=today or datetime.now(TZ).date()
    return date.fromisoformat(str(rule['earliest'])),today+timedelta(days=int(rule['max_days_after_import']))

def parse_time(value):
    # A missing time is never imputed; the caller may only match an existing sample of that day.
    if value is None or str(value).strip()=='': return None,'time missing'
    if isinstance(value,datetime): return value.time(),None
    if isinstance(value,time): return value,None
    try:
        if isinstance(value,(int,float)) and not isinstance(value,bool) and 0<=value<1:
            seconds=round(float(value)*86400)
            return time(seconds//3600,(seconds%3600)//60,seconds%60),None
        return time.fromisoformat(str(value).strip().replace('.',':')),None
    except ValueError as exc: raise RowError(f'not a time: {value!r}') from exc

def local_timestamp(day,clock):
    naive=datetime.combine(day,clock)
    if naive.tzinfo is not None: raise RowError('expected a local time without zone')
    candidates=[]
    for fold in (0,1):
        candidate=naive.replace(tzinfo=TZ,fold=fold)
        if candidate.astimezone(timezone.utc).astimezone(TZ).replace(tzinfo=None)==naive:
            candidates.append(candidate)
    if not candidates: raise RowError('nonexistent Berlin time during DST change')
    if len({x.utcoffset() for x in candidates})>1: raise RowError('ambiguous Berlin time during DST change')
    return candidates[0]

def raw_value(value):
    if isinstance(value,(date,datetime,time)): return {'type':type(value).__name__,'value':value.isoformat()}
    if isinstance(value,float) and not math.isfinite(value): return {'type':'float','value':str(value)}
    return value

def blank(v): return v is None or str(v).strip()==''

def read_records(path,window=None):
    earliest,latest=window or date_window()
    workbook=load_workbook(path,data_only=False,read_only=True)
    records=[]
    try:
        unexpected=set(workbook.sheetnames)-{'Analysen','Dosierung'}
        if unexpected:raise ValueError(f'Unrecognized sheets require review: {sorted(unexpected)}')
        for name in ('Analysen','Dosierung'):
            if name not in workbook.sheetnames: raise ValueError(f'Missing sheet {name}')
            rows=list(workbook[name].iter_rows(values_only=True))
            header_no=next((i for i,row in enumerate(rows) if key(row[0])=='datum'),None)
            if header_no is None: raise ValueError(f'{name}: Datum header missing')
            names=[key(h).split('[')[0].strip() for h in rows[header_no] if h is not None]
            if len(names)!=len(set(names)):raise ValueError(f'{name}: duplicate column headers')
            header={key(h).split('[')[0].strip():i for i,h in enumerate(rows[header_no]) if h is not None}
            required={'datum','uhrzeit','bad'} | (set(PARAMS) if name=='Analysen' else {'zusatz','menge','einheit'})
            if not required<=header.keys(): raise ValueError(f'{name}: missing columns {required-header.keys()}')
            if name=='Analysen':
                for parameter in PARAMS:
                    label=key(rows[header_no][header[parameter]])
                    if re.sub(r'\s+','',label)!=parameter+'[g/l]':
                        raise ValueError(f'{name}: expected g/L units for {parameter}, got {label!r}')
            headed={i for i,h in enumerate(rows[header_no]) if h is not None}
            for no,row in enumerate(rows,1):
                record={'sheet':name,'row':no,'raw':[raw_value(v) for v in row],'status':'loaded','reason':None,'parsed':None}
                records.append(record)
                if no<=header_no+1:
                    record.update(status='ignored',reason='title or header row');continue
                if all(blank(v) for v in row):
                    record.update(status='ignored',reason='blank row');continue
                stray=[i for i,v in enumerate(row) if i not in headed and not blank(v)]
                if stray:
                    record.update(status='quarantined',reason='value in a column without header: '+', '.join(get_column_letter(i+1) for i in stray));continue
                if key(row[0]).startswith('mittelwert'):
                    # A summary line has no time and no tank; anything else is data under a summary label.
                    if blank(row[header['uhrzeit']]) and blank(row[header['bad']]):
                        record.update(status='ignored',reason='weekly summary; not a measurement');continue
                    record.update(status='quarantined',reason='summary label on a row with time or tank; requires review');continue
                ident=None # the sample or daily total this row stands for, once tank, time and additive are readable
                try:
                    if any(isinstance(v,str) and v.startswith('=') for v in row): raise RowError('formula cell requires explicit value review')
                    day=parse_date(row[header['datum']])
                    if not earliest<=day<=latest: raise RowError(f'date {day.isoformat()} outside plausible window {earliest.isoformat()}..{latest.isoformat()}')
                    clock_time,note=parse_time(row[header['uhrzeit']])
                    ts=local_timestamp(day,clock_time) if clock_time is not None else None
                    if ts is not None and ts>datetime.now(timezone.utc): raise RowError(f'time {ts.isoformat()} is later than now')
                    tank=ALIASES.get(key(row[header['bad']]))
                    if tank is None: raise RowError('unknown tank alias')
                    # From here on the row is this tank's row for this date, even if a value turns out unreadable.
                    record['cover']=(tank,day)
                    if name=='Analysen':
                        ident=(tank,ts)
                        values={p:parse_number(row[header[h]]) for h,p in PARAMS.items()}
                        if any(v is not None and v!=v.quantize(Decimal('.0001')) for v in values.values()):
                            raise RowError('concentration exceeds stored precision of 4 decimals')
                        if any(values[p] is None for p in ('ZN','NI','NAOH')): raise RowError('mandatory Zn, Ni or NaOH measurement missing')
                        missing=[p for p,v in values.items() if v is None]
                        record['reason']='; '.join(x for x in (note,'not measured: '+', '.join(missing) if missing else None) if x) or None
                        record['parsed']={'tank':tank,'ts':ts,'day':day,'values':values}
                    else:
                        additive=str(row[header['zusatz']]).strip()
                        if additive not in ('nickel_replenisher','brightener'): raise RowError('unknown additive')
                        ident=(tank,day,additive)
                        amount=parse_number(row[header['menge']]);unit=key(row[header['einheit']])
                        if amount is None or amount<=0 or unit not in ('l','ml'): raise RowError('dose needs positive amount and L/mL unit')
                        ml=amount*(1000 if unit=='l' else 1)
                        if ml!=ml.quantize(Decimal('.001')):raise RowError('dose exceeds stored precision of 0.001 mL')
                        record['reason']=note
                        record['parsed']={'tank':tank,'ts':ts,'day':day,'additive':additive,'ml':ml.quantize(Decimal('.001'))}
                except (RowError,IndexError,InvalidOperation) as exc:
                    record.update(status='quarantined',reason=str(exc),parsed=None,unread=ident)
    finally: workbook.close()
    # In-file identity: an analysis is (tank, timestamp), or (tank, day) while its time is missing;
    # a daily-total dose is (tank, Berlin day, additive). Two dose rows are the same event only when
    # the amounts agree and the times agree or one of them is missing.
    groups=defaultdict(list)
    for r in records:
        if r['parsed']:
            p=r['parsed']
            ident=(p['day'],p['additive']) if r['sheet']=='Dosierung' else (p['ts'] or p['day'],)
            groups[(r['sheet'],p['tank'],*ident)].append(r)
    def same_event(a,b):
        if 'additive' in a: return a['ml']==b['ml'] and (a['ts'] is None or b['ts'] is None or a['ts']==b['ts'])
        return a==b
    for same in groups.values():
        if any(not same_event(a['parsed'],b['parsed']) for a,b in combinations(same,2)):
            amounts={r['parsed'].get('ml') for r in same}
            reason=('one daily total entered at different times; enter one row per day' if 'additive' in same[0]['parsed'] and len(amounts)==1
                    else 'conflicting rows with the same event key')
            for r in same: r.update(status='quarantined',reason=reason,parsed=None)
        else:
            for r in same[1:]:
                exact=r['parsed']==same[0]['parsed']
                r.update(status='duplicate',reason=f'{"exact duplicate" if exact else "same daily total"} of row {same[0]["row"]}',parsed=None)
    # Identical values at two different times on one day look like a changed time or a copied row.
    copies=defaultdict(list)
    for r in records:
        p=r['parsed']
        if p and r['sheet']=='Analysen' and p['ts'] is not None:
            copies[(p['tank'],p['day'],tuple(sorted(p['values'].items())))].append(r)
    for same in copies.values():
        if len(same)>1:
            for r in same: r.update(status='quarantined',reason='identical values at different times on the same day; changed time or copied row?',parsed=None)
    return records

def berlin_day(ts): return ts.astimezone(TZ).date()
def clock(ts): return ts.astimezone(TZ).strftime('%H:%M')

def import_workbook(conn,path,apply=False,window=None,reuse=True):
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    if apply: conn.execute('SELECT pg_advisory_xact_lock(72001001)')
    existing=conn.execute('SELECT report FROM lab.import_file WHERE sha256=%s',(digest,)).fetchone()
    if existing and reuse:
        # The stored report describes the database at that import. Repeat the comparison as a dry run
        # so rows added to or removed from the database since then are visible now.
        now=import_workbook(conn,path,False,window,reuse=False)
        return {**existing[0],'reused':True,'inserted_this_run':0,
                'review_required_now':now['review_required'],'database_only_rows_now':now['database_only_rows']}
    records=read_records(path,window)
    tanks=dict(conn.execute('SELECT code,tank_id FROM ref.tank').fetchall())
    codes={v:k for k,v in tanks.items()}
    reconciliation={'matched_values':0,'source_only_values':0,'database_only_not_measured':0,'conflicting_rows':0,
      'rekeyed_rows':0,'time_missing_matched':0,'retimed_daily_totals':0,'database_only_analyses':0,'database_only_doses':0}
    for record in records:
        p=record['parsed']
        if not p:continue
        tank_id=tanks.get(p['tank'])
        if tank_id is None:
            record.update(status='quarantined',reason='tank missing in database',parsed=None);continue
        p['tank_id']=tank_id
    def day_samples(tank_id,day):
        return [r[0] for r in conn.execute('''SELECT sampled_at FROM lab.analysis WHERE tank_id=%s
          AND (sampled_at AT TIME ZONE 'Europe/Berlin')::date=%s ORDER BY sampled_at''',(tank_id,day)).fetchall()]
    # The workbook covers, per sheet and tank, the Berlin dates from its first to its last row with a
    # readable tank and date (including rows quarantined later as conflicting).
    coverage={}
    for r in records:
        if r.get('cover') and r['cover'][0] in tanks:
            k=(r['sheet'],tanks[r['cover'][0]]);d=r['cover'][1];lo,hi=coverage.get(k,(d,d))
            coverage[k]=(min(lo,d),max(hi,d))
    analyses=[r for r in records if r['parsed'] and r['sheet']=='Analysen']
    explained={}   # database samples already accounted for by a quarantined workbook row
    # A row without a time may only confirm the single existing sample of that tank and Berlin date.
    timed={(r['parsed']['tank_id'],r['parsed']['ts']):r for r in analyses if r['parsed']['ts']}
    for r in analyses:
        p=r['parsed']
        if p['ts'] is not None:continue
        stored=day_samples(p['tank_id'],p['day'])
        if len(stored)!=1:
            for ts in stored: explained[(p['tank_id'],ts)]=r['row']
            r.update(status='quarantined',parsed=None,reason='time missing and '+
              ('no database sample on this date' if not stored else f'{len(stored)} database samples on this date')+
              '; the sampling time must be entered')
            continue
        other=timed.get((p['tank_id'],stored[0]))
        if other:
            if other['parsed'] and other['parsed']['values']==p['values']:
                r.update(status='duplicate',reason=f'exact duplicate of row {other["row"]} (time missing)',parsed=None)
            else:
                r.update(status='quarantined',reason=f'time missing; row {other["row"]} has other values for the only sample on this date',parsed=None)
            continue
        p['ts']=stored[0];reconciliation['time_missing_matched']+=1
        r['reason']=r['reason'].replace('time missing',f'time missing; matched the only database sample of that date ({clock(stored[0])})',1)
    wb_analyses={(r['parsed']['tank_id'],r['parsed']['ts']) for r in analyses if r['parsed']}
    wb_doses={f'{r["parsed"]["tank"]}|{r["parsed"]["day"].isoformat()}|{r["parsed"]["additive"]}'
              for r in records if r['parsed'] and r['sheet']=='Dosierung'}
    # A row quarantined because one of its values is unreadable is still that sample's row in this workbook:
    # its database counterpart must not also be reported as missing from the workbook.
    for r in records:
        if not r.get('unread') or r['unread'][0] not in tanks:continue
        tank,when=r['unread'][:2]
        if r['sheet']=='Dosierung':wb_doses.add(f'{tank}|{when.isoformat()}|{r["unread"][2]}')
        elif when is not None:wb_analyses.add((tanks[tank],when))
        else:
            for ts in day_samples(tanks[tank],r['cover'][1]):explained[(tanks[tank],ts)]=r['row']
    pending=[]
    for record in records:
        p=record['parsed']
        if not p:continue
        if record['sheet']=='Analysen':
            stored=day_samples(p['tank_id'],berlin_day(p['ts']))
            if p['ts'] not in stored:
                # Same tank and date, another time, and that sample has no workbook row: a re-keyed correction.
                others=[ts for ts in stored if (p['tank_id'],ts) not in wb_analyses]
                if others:
                    for ts in others: explained[(p['tank_id'],ts)]=record['row']
                    record.update(status='quarantined',parsed=None,reason=f'changed time? database sample of {berlin_day(p["ts"]).isoformat()} at '
                      +', '.join(clock(ts) for ts in others)+f' has no workbook row, workbook says {clock(p["ts"])}; review before importing')
                    reconciliation['rekeyed_rows']+=1;continue
                # A new time whose values equal a stored sample of the same day: changed time or copied row.
                twin=next((ts for ts in stored if (lambda vals:vals and all(vals.get(k)==v for k,v in p['values'].items() if v is not None))(
                  dict(conn.execute('''SELECT r.parameter_code,r.value FROM lab.analysis a JOIN lab.analysis_result r USING(analysis_id)
                  WHERE a.tank_id=%s AND a.sampled_at=%s''',(p['tank_id'],ts)).fetchall()))),None)
                if twin:
                    record.update(status='quarantined',parsed=None,reason=f'same values as the database sample of {berlin_day(twin).isoformat()} at '
                      f'{clock(twin)}, workbook says {clock(p["ts"])}; changed time or copied row? review before importing')
                    reconciliation['rekeyed_rows']+=1;continue
            old=dict(conn.execute('''SELECT r.parameter_code,r.value FROM lab.analysis a JOIN lab.analysis_result r USING(analysis_id)
              WHERE a.tank_id=%s AND a.sampled_at=%s''',(p['tank_id'],p['ts'])).fetchall())
            conflict=[k for k,v in p['values'].items() if v is not None and k in old and v.quantize(Decimal('.0001'))!=old[k]]
            if conflict:
                record.update(status='quarantined',reason='database conflict: '+', '.join(conflict),parsed=None)
                reconciliation['conflicting_rows']+=1;continue
            for k,v in p['values'].items():
                if v is None:
                    if k in old:reconciliation['database_only_not_measured']+=1
                elif k in old:reconciliation['matched_values']+=1
                else:reconciliation['source_only_values']+=1
            p['new']=p['ts'] not in stored
            p['source_only']=sum(v is not None and k not in old for k,v in p['values'].items())
        else:
            p['event_key']=f'{p["tank"]}|{p["day"].isoformat()}|{p["additive"]}'
            old=conn.execute('SELECT amount,unit,dosed_at FROM lab.dosing WHERE event_key=%s',(p['event_key'],)).fetchone()
            if old and (old[1]!='mL' or old[0]!=p['ml']):
                record.update(status='quarantined',reason='database conflict: daily dose changed',parsed=None)
                reconciliation['conflicting_rows']+=1;continue
            if old is None and p['ts'] is None:
                record.update(status='quarantined',reason='time missing on a new daily total; the dosing time must be entered',parsed=None);continue
            p['new']=old is None
            if old and p['ts'] is None:
                record['reason']=f'time missing; same daily total as the database entry at {clock(old[2])}';p['ts']=old[2]
            elif old and p['ts']!=old[2]:
                record['reason']=f'same daily total, stored at {clock(old[2])}, workbook says {clock(p["ts"])}; stored time kept'
                reconciliation['retimed_daily_totals']+=1
        pending.append(record)
    # Database -> workbook: rows inside the covered tanks and dates that this workbook does not contain.
    database_only=[]
    for (sheet,tank_id),(lo,hi) in sorted(coverage.items()):
        if sheet=='Analysen':
            for aid,ts in conn.execute('''SELECT analysis_id,sampled_at FROM lab.analysis WHERE tank_id=%s
              AND (sampled_at AT TIME ZONE 'Europe/Berlin')::date BETWEEN %s AND %s AND NOT sampled_at=ANY(%s)
              ORDER BY sampled_at''',(tank_id,lo,hi,[ts for t,ts in wb_analyses if t==tank_id])).fetchall():
                if (tank_id,ts) in explained:continue
                # A new row carrying exactly this sample's values is the same sample with a changed date.
                stored=dict(conn.execute('SELECT parameter_code,value FROM lab.analysis_result WHERE analysis_id=%s',(aid,)).fetchall())
                twin=next((r for r in pending if r['sheet']=='Analysen' and r['parsed']['tank_id']==tank_id and r['parsed']['new']
                  and all(stored.get(k)==v for k,v in r['parsed']['values'].items() if v is not None)),None)
                if twin:
                    pending.remove(twin);reconciliation['rekeyed_rows']+=1
                    reconciliation['source_only_values']-=twin['parsed']['source_only']
                    twin.update(status='quarantined',parsed=None,reason=f'changed date? same values as the database sample of '
                      f'{berlin_day(ts).isoformat()} at {clock(ts)}, which has no workbook row; review before importing')
                    continue
                database_only.append({'sheet':sheet,'tank':codes[tank_id],'analysis_id':aid,'sampled_at':ts.astimezone(TZ).isoformat()})
                reconciliation['database_only_analyses']+=1
        else:
            for did,ts,additive,amount,unit in conn.execute('''SELECT dosing_id,dosed_at,additive,amount,unit FROM lab.dosing
              WHERE tank_id=%s AND (dosed_at AT TIME ZONE 'Europe/Berlin')::date BETWEEN %s AND %s
              AND NOT coalesce(event_key,'')=ANY(%s) ORDER BY dosed_at,additive''',(tank_id,lo,hi,sorted(wb_doses))).fetchall():
                twin=next((r for r in pending if r['sheet']=='Dosierung' and r['parsed']['tank_id']==tank_id and r['parsed']['new']
                  and r['parsed']['additive']==additive and unit=='mL' and r['parsed']['ml']==amount),None)
                if twin:
                    pending.remove(twin);reconciliation['rekeyed_rows']+=1
                    twin.update(status='quarantined',parsed=None,reason=f'changed date? same amount as the database {additive} total of '
                      f'{berlin_day(ts).isoformat()}, which has no workbook row; review before importing')
                    continue
                database_only.append({'sheet':sheet,'tank':codes[tank_id],'dosing_id':did,'dosed_at':ts.astimezone(TZ).isoformat(),
                  'additive':additive,'amount':str(amount),'unit':unit})
                reconciliation['database_only_doses']+=1
    inserted={'analyses':0,'results':0,'doses':0}
    if apply:
        # Sort new analyses so timeline changes are deterministic.
        for record in sorted(pending,key=lambda r:r['parsed']['ts']):
            p=record['parsed']
            if record['sheet']=='Analysen':
                cur=conn.execute('''INSERT INTO lab.analysis(tank_id,sampled_at,source) VALUES(%s,%s,%s)
                  ON CONFLICT(tank_id,sampled_at) DO NOTHING''',(p['tank_id'],p['ts'],f'excel:{digest}#{record["row"]}'))
                inserted['analyses']+=cur.rowcount
                aid=conn.execute('SELECT analysis_id FROM lab.analysis WHERE tank_id=%s AND sampled_at=%s',(p['tank_id'],p['ts'])).fetchone()[0]
                for param,value in p['values'].items():
                    if value is not None:
                        cur=conn.execute('INSERT INTO lab.analysis_result VALUES(%s,%s,%s,%s) ON CONFLICT DO NOTHING',(aid,param,value,METHODS[param]))
                        inserted['results']+=cur.rowcount
            else:
                cur=conn.execute('''INSERT INTO lab.dosing(tank_id,dosed_at,additive,amount,unit,reason,source,event_key)
                  VALUES(%s,%s,%s,%s,'mL','Imported daily total',%s,%s) ON CONFLICT(event_key) DO NOTHING''',
                  (p['tank_id'],p['ts'],p['additive'],p['ml'],f'excel:{digest}#{record["row"]}',p['event_key']))
                inserted['doses']+=cur.rowcount
    issues=[{'sheet':r['sheet'],'row':r['row'],'status':r['status'],'reason':r['reason']} for r in records if r['reason']]
    for d in database_only:
        what=f'analysis {d["analysis_id"]} at {d["sampled_at"]}' if 'analysis_id' in d else f'{d["additive"]} dose {d["dosing_id"]} at {d["dosed_at"]}'
        issues.append({'sheet':d['sheet'],'row':None,'status':'database_only','reason':f'database {what} has no row in this workbook version'})
    report={'file':path.name,'sha256':digest,'applied':apply,'reused':False,'rows':len(records),
      'statuses':dict(Counter(r['status'] for r in records)),
      'sheets':{name:dict(Counter(r['status'] for r in records if r['sheet']==name)) for name in ('Analysen','Dosierung')},
      'reconciliation_before':reconciliation,'inserted':inserted,'inserted_this_run':sum(inserted.values()),
      'unexplained_rows':sum(r['status'] in ('quarantined','ignored','duplicate') and not r['reason'] for r in records),
      'review_required':{'quarantined_rows':sum(r['status']=='quarantined' for r in records),
        'database_only_rows':len(database_only),'retimed_daily_totals':reconciliation['retimed_daily_totals']},
      'coverage':{f'{sheet}|{codes[t]}':[lo.isoformat(),hi.isoformat()] for (sheet,t),(lo,hi) in sorted(coverage.items())},
      'database_only_rows':database_only,'issues':issues}
    if apply:
        after={'source_only_values':0,'different_values':0}
        for record in pending:
            if record['sheet']!='Analysen':continue
            p=record['parsed']
            stored=dict(conn.execute('''SELECT r.parameter_code,r.value FROM lab.analysis a JOIN lab.analysis_result r USING(analysis_id)
              WHERE a.tank_id=%s AND a.sampled_at=%s''',(p['tank_id'],p['ts'])).fetchall())
            for param,value in p['values'].items():
                if value is None:continue
                if param not in stored:after['source_only_values']+=1
                elif stored[param]!=value:after['different_values']+=1
        report['reconciliation_after']=after
        conn.execute('INSERT INTO lab.import_file(sha256,file_name,report) VALUES(%s,%s,%s)',(digest,path.name,Jsonb(report)))
        with conn.cursor() as cur:
            cur.executemany('''INSERT INTO lab.stg_excel_row(file_name,sheet_name,row_no,raw,status,reason,import_hash)
              VALUES(%s,%s,%s,%s,%s,%s,%s)''',[(path.name,r['sheet'],r['row'],Jsonb(r['raw']),r['status'],r['reason'],digest) for r in records])
    return report

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook',nargs='?',type=__import__('pathlib').Path,default=ROOT/'data/raw/lab/Laborjournal_ZnNi_L1.xlsx')
    parser.add_argument('--apply',action='store_true')
    args=parser.parse_args()
    with connect('lab_etl') as conn: report=import_workbook(conn,args.workbook,args.apply)
    out=ROOT/'docs/results'/('etl_applied.json' if args.apply else 'etl_dry_run.json')
    out.write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k not in ('issues','database_only_rows')},indent=2))

if __name__=='__main__': main()
